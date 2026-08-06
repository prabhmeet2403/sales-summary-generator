"""
tests/fixture_builder.py
========================
Shared helper for building minimal, purpose-built synthetic input
workbooks for the Track 2 test suite. Not itself a test -- imported by
the test files in this directory.
"""
from __future__ import annotations

import openpyxl
from datetime import datetime
from pathlib import Path

MONTHS = list(range(1, 13))


def new_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales by Customer- 2026"

    col = 1
    ws.cell(row=2, column=col, value="Name"); col += 1
    ws.cell(row=2, column=col, value="POC"); col += 1
    ws.cell(row=2, column=col, value="Service"); col += 1

    month_cols = {}
    for m in MONTHS:
        rev_c, sal_c, marg_c = col, col + 1, col + 2
        month_cols[m] = (rev_c, sal_c, marg_c)
        d = datetime(2026, m, 1)
        ws.cell(row=2, column=rev_c, value=d); ws.cell(row=1, column=rev_c, value="Actual")
        ws.cell(row=2, column=sal_c, value=d); ws.cell(row=1, column=sal_c, value="Salary")
        ws.cell(row=2, column=marg_c, value=d); ws.cell(row=1, column=marg_c, value="Margin")
        col += 3

    total_rev_col = col; ws.cell(row=2, column=col, value="Total Revenue"); col += 1
    total_marg_col = col; ws.cell(row=2, column=col, value="Total Margin"); col += 1
    comments_col = col; ws.cell(row=2, column=col, value="Comments"); col += 1
    confidence_col = col; ws.cell(row=2, column=col, value="Renewal Confidence"); col += 1
    group_col = col; ws.cell(row=2, column=col, value="Group"); col += 1
    subgroup_col = col; ws.cell(row=2, column=col, value="Sub-Group"); col += 1

    layout = {
        "month_cols": month_cols, "total_rev_col": total_rev_col, "total_marg_col": total_marg_col,
        "comments_col": comments_col, "confidence_col": confidence_col,
        "group_col": group_col, "subgroup_col": subgroup_col,
    }
    return wb, ws, layout


def write_row(ws, layout, r, name, poc, service, group, subgroup, revenue_month, revenue_val, margin_val):
    ws.cell(row=r, column=1, value=name)
    ws.cell(row=r, column=2, value=poc)
    ws.cell(row=r, column=3, value=service)
    for m in MONTHS:
        rev_c, sal_c, marg_c = layout["month_cols"][m]
        if m == revenue_month:
            ws.cell(row=r, column=rev_c, value=revenue_val)
            ws.cell(row=r, column=marg_c, value=margin_val)
            ws.cell(row=r, column=sal_c, value=revenue_val - margin_val)
        else:
            ws.cell(row=r, column=rev_c, value=0)
            ws.cell(row=r, column=marg_c, value=0)
            ws.cell(row=r, column=sal_c, value=0)
    ws.cell(row=r, column=layout["total_rev_col"], value=revenue_val)
    ws.cell(row=r, column=layout["total_marg_col"], value=margin_val)
    ws.cell(row=r, column=layout["group_col"], value=group)
    ws.cell(row=r, column=layout["subgroup_col"], value=subgroup)


def add_section(ws, layout, r, heading, rows_data):
    """rows_data: list of dicts with keys name,poc,service,group,subgroup,month,revenue,margin"""
    ws.cell(row=r, column=1, value=heading)
    r += 2
    for row_data in rows_data:
        write_row(ws, layout, r, row_data["name"], row_data.get("poc"), row_data.get("service", "Project"),
                   row_data["group"], row_data["subgroup"], row_data["month"], row_data["revenue"], row_data["margin"])
        r += 1
    ws.cell(row=r, column=1, value=f"Sub-total : {heading.split('- ')[-1] if '- ' in heading else heading}")
    r += 2
    return r
