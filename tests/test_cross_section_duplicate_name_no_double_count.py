"""
tests/test_cross_section_duplicate_name_no_double_count.py
=============================================================
Verifies requirement 1 and 2 together, under the hardest condition:
the SAME customer name appears as a genuine, separate row in BOTH
Track 1 and Track 2, each with its own distinct real revenue. Neither
row may be double-counted into the other section, and each section's
total must reflect only its own row.

Usage:
    python tests/test_cross_section_duplicate_name_no_double_count.py
"""
from __future__ import annotations
import subprocess
import sys
from pathlib import Path

TEST_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TEST_DIR))
sys.path.insert(0, str(TEST_DIR.parent))
from fixture_builder import new_workbook, write_row  # noqa: E402

import main as sfae_main  # noqa: E402
from openpyxl import load_workbook

TRACK1_REVENUE = 555001
TRACK2_REVENUE = 777003


def build_fixture(path: Path):
    wb, ws, layout = new_workbook()
    r = 4
    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 1"); ws.cell(row=r, column=44, value="Track 1-Secured"); r += 2
    write_row(ws, layout, r, "Shared Name Co - Track 1 engagement", "Priya", "Project", "Shared Name Co", "DS10_Secured", 1, TRACK1_REVENUE, 55500); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 1"); r += 2

    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 2"); ws.cell(row=r, column=44, value="Track 2-Secured"); r += 2
    write_row(ws, layout, r, "Shared Name Co - Track 2 engagement", "Dana", "Staff Aug", "Shared Name Co", "DS20_Secured", 1, TRACK2_REVENUE, 77700); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 2"); r += 1

    wb.save(path)


def main() -> int:
    fixture = TEST_DIR / "fixtures" / "test10_cross_section_duplicate_name.xlsx"
    build_fixture(fixture)

    out_dir = TEST_DIR.parent / "output" / "_test10"
    out_dir.mkdir(parents=True, exist_ok=True)
    rc = sfae_main.main(["--input", str(fixture), "--output-dir", str(out_dir), "--year", "2026"])
    if rc != 0:
        print("FAIL - pipeline did not exit cleanly")
        return 1

    generated = list(out_dir.glob("Sales_and_Forecast_Summary_*.xlsx"))[0]
    recalc_dir = TEST_DIR.parent / "output" / "_test10_recalc"
    recalc_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
         "--outdir", str(recalc_dir), str(generated)],
        capture_output=True, timeout=60,
    )
    recalculated = recalc_dir / generated.name

    wb = load_workbook(recalculated, data_only=True)
    ws2 = wb["2026 Monthly Performance"]

    def all_appearances(name):
        current = None
        found = []
        for r in range(1, ws2.max_row + 1):
            v = ws2.cell(row=r, column=1).value
            if v in ("Solutions and Staff Augmentation (Projects) - Track 1",
                      "Solutions and Staff Augmentation (Projects) - Track 2"):
                current = v
                continue
            if v == name:
                monthly = [ws2.cell(row=r, column=c).value for c in range(3, 27, 2)]
                total = sum(x for x in monthly if isinstance(x, (int, float)))
                found.append((current, total))
        return found

    appearances = all_appearances("Shared Name Co")
    failures = []
    if len(appearances) != 2:
        failures.append(f"Expected exactly 2 rows named 'Shared Name Co', found {len(appearances)}: {appearances}")
    else:
        by_section = dict(appearances)
        if by_section.get("Solutions and Staff Augmentation (Projects) - Track 1") != TRACK1_REVENUE:
            failures.append(f"Track 1 instance: expected {TRACK1_REVENUE}, got {by_section.get('Solutions and Staff Augmentation (Projects) - Track 1')}")
        if by_section.get("Solutions and Staff Augmentation (Projects) - Track 2") != TRACK2_REVENUE:
            failures.append(f"Track 2 instance: expected {TRACK2_REVENUE}, got {by_section.get('Solutions and Staff Augmentation (Projects) - Track 2')}")

    def get_subtotal(label):
        for r in range(1, ws2.max_row + 1):
            if ws2.cell(row=r, column=1).value == label:
                monthly = [ws2.cell(row=r, column=c).value for c in range(3, 27, 2)]
                return sum(x for x in monthly if isinstance(x, (int, float)))
        return None

    t1_sub = get_subtotal("Subtotal : Track 1")
    t2_sub = get_subtotal("Subtotal : Track 2")
    if t1_sub != TRACK1_REVENUE:
        failures.append(f"Track 1 subtotal={t1_sub}, expected exactly {TRACK1_REVENUE} (no contamination from the Track 2 instance)")
    if t2_sub != TRACK2_REVENUE:
        failures.append(f"Track 2 subtotal={t2_sub}, expected exactly {TRACK2_REVENUE} (no contamination from the Track 1 instance)")

    if failures:
        print("FAIL:")
        for f in failures:
            print(" -", f)
        return 1
    print(f"PASS - 'Shared Name Co' correctly kept as two separate rows: Track 1={t1_sub}, Track 2={t2_sub}, no double-counting")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
