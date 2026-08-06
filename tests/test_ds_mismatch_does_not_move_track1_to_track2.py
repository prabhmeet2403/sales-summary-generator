"""
tests/test_ds_mismatch_does_not_move_track1_to_track2.py
==========================================================
Verifies requirement 7 (direction 1): a row physically inside Track 1,
tagged with Track 2's Sub-Group code (DS20_Secured) instead of Track
1's own (DS10_Secured), must still be counted in Track 1 -- and must
NOT appear in, or contribute to, Track 2.

Usage:
    python tests/test_ds_mismatch_does_not_move_track1_to_track2.py
"""
from __future__ import annotations
import subprocess
import sys
from pathlib import Path

TEST_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TEST_DIR))
sys.path.insert(0, str(TEST_DIR.parent))
from fixture_builder import new_workbook, write_row  # noqa: E402

import main as sfae_main  # noqa: E402
from openpyxl import load_workbook

EXPECTED_REVENUE = 137931


def build_fixture(path: Path):
    wb, ws, layout = new_workbook()
    r = 4
    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 1"); ws.cell(row=r, column=44, value="Track 1-Secured"); r += 2
    write_row(ws, layout, r, "Alpha Mismatch Co - wrong DS code", "Priya", "Project", "Alpha Mismatch Co", "DS20_Secured", 1, EXPECTED_REVENUE, 13793); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 1"); r += 2

    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 2"); ws.cell(row=r, column=44, value="Track 2-Secured"); r += 2
    write_row(ws, layout, r, "Bravo Real Co - Track 2", "Dana", "Staff Aug", "Bravo Real Co", "DS20_Secured", 1, 246801, 24680); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 2"); r += 1

    wb.save(path)


def main() -> int:
    fixture = TEST_DIR / "fixtures" / "test8_ds_mismatch_t1_to_t2.xlsx"
    build_fixture(fixture)

    out_dir = TEST_DIR.parent / "output" / "_test8"
    out_dir.mkdir(parents=True, exist_ok=True)
    rc = sfae_main.main(["--input", str(fixture), "--output-dir", str(out_dir), "--year", "2026"])
    if rc != 0:
        print("FAIL - pipeline did not exit cleanly")
        return 1

    generated = list(out_dir.glob("Sales_and_Forecast_Summary_*.xlsx"))[0]
    recalc_dir = TEST_DIR.parent / "output" / "_test8_recalc"
    recalc_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
         "--outdir", str(recalc_dir), str(generated)],
        capture_output=True, timeout=60,
    )
    recalculated = recalc_dir / generated.name

    wb = load_workbook(recalculated, data_only=True)
    ws2 = wb["2026 Monthly Performance"]

    def section_of(name):
        current = None
        for r in range(1, ws2.max_row + 1):
            v = ws2.cell(row=r, column=1).value
            if v in ("Solutions and Staff Augmentation (Projects) - Track 1",
                      "Solutions and Staff Augmentation (Projects) - Track 2"):
                current = v
                continue
            if v == name:
                monthly = [ws2.cell(row=r, column=c).value for c in range(3, 27, 2)]
                total = sum(x for x in monthly if isinstance(x, (int, float)))
                return current, total
        return None, None

    section, total = section_of("Alpha Mismatch Co")
    failures = []
    if section != "Solutions and Staff Augmentation (Projects) - Track 1":
        failures.append(f"Alpha Mismatch Co found in {section!r}, expected Track 1")
    if total != EXPECTED_REVENUE:
        failures.append(f"Alpha Mismatch Co total={total}, expected {EXPECTED_REVENUE}")

    t2_subtotal = None
    for r in range(1, ws2.max_row + 1):
        if ws2.cell(row=r, column=1).value == "Subtotal : Track 2":
            monthly = [ws2.cell(row=r, column=c).value for c in range(3, 27, 2)]
            t2_subtotal = sum(x for x in monthly if isinstance(x, (int, float)))
    if t2_subtotal != 246801:
        failures.append(f"Track 2 subtotal={t2_subtotal}, expected exactly 246801 (Bravo Real Co only -- Alpha must not have leaked in)")

    if failures:
        print("FAIL:")
        for f in failures:
            print(" -", f)
        return 1
    print(f"PASS - Alpha Mismatch Co stays in Track 1 ({total}), Track 2 subtotal unaffected ({t2_subtotal})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
