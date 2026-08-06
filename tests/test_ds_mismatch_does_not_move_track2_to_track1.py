"""
tests/test_ds_mismatch_does_not_move_track2_to_track1.py
==========================================================
Verifies requirement 7 (direction 2, the mirror of test 8): a row
physically inside Track 2, tagged with Track 1's Sub-Group code
(DS10_Secured) instead of Track 2's own (DS20_Secured), must still be
counted in Track 2 -- and must NOT appear in, or contribute to, Track
1. This is the exact scenario (a Track 1 row copy-pasted into Track 2
without updating Sub-Group) that originally motivated the
section_key-based fix.

Usage:
    python tests/test_ds_mismatch_does_not_move_track2_to_track1.py
"""
from __future__ import annotations
import subprocess
import sys
from pathlib import Path

TEST_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TEST_DIR))
sys.path.insert(0, str(TEST_DIR.parent))
from fixture_builder import new_workbook, write_row  # noqa: E402

import main as sfae_main  # noqa: E402
from openpyxl import load_workbook

EXPECTED_REVENUE = 913579


def build_fixture(path: Path):
    wb, ws, layout = new_workbook()
    r = 4
    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 1"); ws.cell(row=r, column=44, value="Track 1-Secured"); r += 2
    write_row(ws, layout, r, "Charlie Real Co - Track 1", "Priya", "Project", "Charlie Real Co", "DS10_Secured", 1, 864209, 86420); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 1"); r += 2

    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 2"); ws.cell(row=r, column=44, value="Track 2-Secured"); r += 2
    write_row(ws, layout, r, "Delta Mismatch Co - copied without updating Sub-Group", "Dana", "Staff Aug", "Delta Mismatch Co", "DS10_Secured", 1, EXPECTED_REVENUE, 91357); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 2"); r += 1

    wb.save(path)


def main() -> int:
    fixture = TEST_DIR / "fixtures" / "test9_ds_mismatch_t2_to_t1.xlsx"
    build_fixture(fixture)

    out_dir = TEST_DIR.parent / "output" / "_test9"
    out_dir.mkdir(parents=True, exist_ok=True)
    rc = sfae_main.main(["--input", str(fixture), "--output-dir", str(out_dir), "--year", "2026"])
    if rc != 0:
        print("FAIL - pipeline did not exit cleanly")
        return 1

    generated = list(out_dir.glob("Sales_and_Forecast_Summary_*.xlsx"))[0]
    recalc_dir = TEST_DIR.parent / "output" / "_test9_recalc"
    recalc_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
         "--outdir", str(recalc_dir), str(generated)],
        capture_output=True, timeout=60,
    )
    recalculated = recalc_dir / generated.name

    wb = load_workbook(recalculated, data_only=True)
    ws2 = wb["2026 Monthly Performance"]

    def section_of(name):
        current = None
        for r in range(1, ws2.max_row + 1):
            v = ws2.cell(row=r, column=1).value
            if v in ("Solutions and Staff Augmentation (Projects) - Track 1",
                      "Solutions and Staff Augmentation (Projects) - Track 2"):
                current = v
                continue
            if v == name:
                monthly = [ws2.cell(row=r, column=c).value for c in range(3, 27, 2)]
                total = sum(x for x in monthly if isinstance(x, (int, float)))
                return current, total
        return None, None

    section, total = section_of("Delta Mismatch Co")
    failures = []
    if section != "Solutions and Staff Augmentation (Projects) - Track 2":
        failures.append(f"Delta Mismatch Co found in {section!r}, expected Track 2")
    if total != EXPECTED_REVENUE:
        failures.append(f"Delta Mismatch Co total={total}, expected {EXPECTED_REVENUE} (this is the exact double-count/drop defect the fix addresses)")

    t1_subtotal = None
    for r in range(1, ws2.max_row + 1):
        if ws2.cell(row=r, column=1).value == "Subtotal : Track 1":
            monthly = [ws2.cell(row=r, column=c).value for c in range(3, 27, 2)]
            t1_subtotal = sum(x for x in monthly if isinstance(x, (int, float)))
    if t1_subtotal != 864209:
        failures.append(f"Track 1 subtotal={t1_subtotal}, expected exactly 864209 (Charlie Real Co only -- Delta must not have leaked in)")

    if failures:
        print("FAIL:")
        for f in failures:
            print(" -", f)
        return 1
    print(f"PASS - Delta Mismatch Co stays in Track 2 ({total}), Track 1 subtotal unaffected ({t1_subtotal})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
