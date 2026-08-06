"""
tests/test_investments_excluded_from_total_secured.py
=======================================================
Verifies requirement 5: Investments' revenue must never be included in
TOTAL Secured, even though Investments shares Sub-Group DS10_Secured
with Track 1 -- proven by giving the Investments row a large,
distinctive, non-zero revenue value and confirming TOTAL Secured does
NOT include it.

Usage:
    python tests/test_investments_excluded_from_total_secured.py
"""
from __future__ import annotations
import subprocess
import sys
from pathlib import Path

TEST_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TEST_DIR))
sys.path.insert(0, str(TEST_DIR.parent))
from fixture_builder import new_workbook, write_row  # noqa: E402

import main as sfae_main  # noqa: E402
from openpyxl import load_workbook

INVESTMENTS_REVENUE = 9999999  # deliberately large and distinctive


def build_fixture(path: Path):
    wb, ws, layout = new_workbook()
    r = 4
    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 1"); ws.cell(row=r, column=44, value="Track 1-Secured"); r += 2
    write_row(ws, layout, r, "Romeo Corp - SOW", "Priya", "Project", "Romeo Corp", "DS10_Secured", 1, 100007, 10007); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 1"); r += 2

    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 2"); ws.cell(row=r, column=44, value="Track 2-Secured"); r += 2
    write_row(ws, layout, r, "Sierra Corp - SOW", "Dana", "Staff Aug", "Sierra Corp", "DS20_Secured", 1, 200007, 20007); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 2"); r += 2

    ws.cell(row=r, column=1, value="TOTAL Secured"); r += 2

    ws.cell(row=r, column=1, value="Staffing"); ws.cell(row=r, column=44, value="Staffing-Secured"); r += 2
    write_row(ws, layout, r, "Tango Corp - Staffing", "Wendy", "Staffing", "Tango Corp", "DS70_Secured", 1, 300007, 30007); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Staffing"); r += 2

    ws.cell(row=r, column=1, value="Investments"); ws.cell(row=r, column=44, value="Investments"); r += 2
    write_row(ws, layout, r, "Uniform Corp - Internal (large)", "Priya", "Internal", "Uniform Corp", "DS10_Secured", 1, INVESTMENTS_REVENUE, 999999); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Investments"); r += 1

    wb.save(path)


def main() -> int:
    fixture = TEST_DIR / "fixtures" / "test5_investments_excluded.xlsx"
    build_fixture(fixture)

    out_dir = TEST_DIR.parent / "output" / "_test5"
    out_dir.mkdir(parents=True, exist_ok=True)
    rc = sfae_main.main(["--input", str(fixture), "--output-dir", str(out_dir), "--year", "2026"])
    if rc != 0:
        print("FAIL - pipeline did not exit cleanly")
        return 1

    generated = list(out_dir.glob("Sales_and_Forecast_Summary_*.xlsx"))[0]
    recalc_dir = TEST_DIR.parent / "output" / "_test5_recalc"
    recalc_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
         "--outdir", str(recalc_dir), str(generated)],
        capture_output=True, timeout=60,
    )
    recalculated = recalc_dir / generated.name

    wb = load_workbook(recalculated, data_only=True)
    ws2 = wb["2026 Monthly Performance"]

    def get_val(label, col=27):
        for r in range(1, ws2.max_row + 1):
            if ws2.cell(row=r, column=1).value == label:
                return ws2.cell(row=r, column=col).value
        return None

    total_secured = get_val("TOTAL Secured")
    investments_subtotal = get_val("Subtotal : Investments")

    if investments_subtotal != INVESTMENTS_REVENUE:
        print(f"FAIL - Investments subtotal itself is wrong: expected {INVESTMENTS_REVENUE}, got {investments_subtotal}")
        return 1

    if total_secured >= INVESTMENTS_REVENUE:
        print(f"FAIL - TOTAL Secured ({total_secured}) appears to include the {INVESTMENTS_REVENUE} Investments figure")
        return 1

    print(f"PASS - Investments subtotal={investments_subtotal} correctly excluded from TOTAL Secured={total_secured}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
