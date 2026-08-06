"""
tests/test_monthly_view_matches_aggregator_ds_mismatch.py
===========================================================
Verifies requirement 6 (the specific, real-world case): a row
physically inside Staffing, tagged with an unmapped Sub-Group code
(DS90_Secured, not in Staffing's configured [70,80,85] -- the real
JSG/Mastek/MatchPoint Solutions/Maxonic/Technology Partners/VDart
case from the real production workbook) must show the SAME non-zero
value on Worksheet 1 and Worksheet 2, not real-vs-zero.

Usage:
    python tests/test_monthly_view_matches_aggregator_ds_mismatch.py
"""
from __future__ import annotations
import subprocess
import sys
from pathlib import Path

TEST_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TEST_DIR))
sys.path.insert(0, str(TEST_DIR.parent))
from fixture_builder import new_workbook, write_row  # noqa: E402

import main as sfae_main  # noqa: E402
from openpyxl import load_workbook

EXPECTED_REVENUE = 468401  # distinctive, unique


def build_fixture(path: Path):
    wb, ws, layout = new_workbook()
    r = 4
    ws.cell(row=r, column=1, value="Staffing"); ws.cell(row=r, column=44, value="Staffing-Secured"); r += 2
    write_row(ws, layout, r, "Zulu Staffing Co - DS90", "Wendy", "Staffing", "Zulu Staffing Co", "DS90_Secured", 1, EXPECTED_REVENUE, 46840); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Staffing"); r += 1

    wb.save(path)


def main() -> int:
    fixture = TEST_DIR / "fixtures" / "test7_ds_mismatch_staffing.xlsx"
    build_fixture(fixture)

    out_dir = TEST_DIR.parent / "output" / "_test7"
    out_dir.mkdir(parents=True, exist_ok=True)
    rc = sfae_main.main(["--input", str(fixture), "--output-dir", str(out_dir), "--year", "2026"])
    if rc != 0:
        print("FAIL - pipeline did not exit cleanly")
        return 1

    generated = list(out_dir.glob("Sales_and_Forecast_Summary_*.xlsx"))[0]
    recalc_dir = TEST_DIR.parent / "output" / "_test7_recalc"
    recalc_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
         "--outdir", str(recalc_dir), str(generated)],
        capture_output=True, timeout=60,
    )
    recalculated = recalc_dir / generated.name

    wb = load_workbook(recalculated, data_only=True)
    ws1, ws2 = wb["Multi-Year Revenue & Margin"], wb["2026 Monthly Performance"]

    def ws1_total(name):
        for r in range(1, ws1.max_row + 1):
            if ws1.cell(row=r, column=1).value == name:
                q = [ws1.cell(row=r, column=c).value for c in (6, 8, 10, 12)]
                return sum(x for x in q if isinstance(x, (int, float)))
        return None

    def ws2_total(name):
        for r in range(1, ws2.max_row + 1):
            if ws2.cell(row=r, column=1).value == name:
                monthly = [ws2.cell(row=r, column=c).value for c in range(3, 27, 2)]
                return sum(x for x in monthly if isinstance(x, (int, float)))
        return None

    v1 = ws1_total("Zulu Staffing Co")
    v2 = ws2_total("Zulu Staffing Co")

    if v1 != EXPECTED_REVENUE:
        print(f"FAIL - Worksheet 1 shows {v1}, expected {EXPECTED_REVENUE}")
        return 1
    if v2 != EXPECTED_REVENUE:
        print(f"FAIL - Worksheet 2 shows {v2}, expected {EXPECTED_REVENUE} (this is the exact defect this test guards against)")
        return 1

    print(f"PASS - Worksheet 1 == Worksheet 2 == {v1}, despite DS90 not matching staffing_secured's configured codes")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
