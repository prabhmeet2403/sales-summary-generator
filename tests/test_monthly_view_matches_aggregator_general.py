"""
tests/test_monthly_view_matches_aggregator_general.py
=======================================================
Verifies requirement 6 (general case): for every group in every
section, Worksheet 1 (built from aggregate_section) and Worksheet 2
(built from monthly_view) must show the same annual total, when
section_key and ds_code agree (the well-formed case).

Usage:
    python tests/test_monthly_view_matches_aggregator_general.py
"""
from __future__ import annotations
import subprocess
import sys
from pathlib import Path

TEST_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TEST_DIR))
sys.path.insert(0, str(TEST_DIR.parent))
from fixture_builder import new_workbook, write_row  # noqa: E402

import main as sfae_main  # noqa: E402
from openpyxl import load_workbook


def build_fixture(path: Path):
    wb, ws, layout = new_workbook()
    r = 4
    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 1"); ws.cell(row=r, column=44, value="Track 1-Secured"); r += 2
    write_row(ws, layout, r, "Victor Corp - SOW", "Priya", "Project", "Victor Corp", "DS10_Secured", 3, 100011, 10011); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 1"); r += 2

    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 2"); ws.cell(row=r, column=44, value="Track 2-Secured"); r += 2
    write_row(ws, layout, r, "Whiskey Corp - SOW", "Dana", "Staff Aug", "Whiskey Corp", "DS20_Secured", 6, 200011, 20011); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 2"); r += 2

    ws.cell(row=r, column=1, value="TOTAL Secured"); r += 2

    ws.cell(row=r, column=1, value="Staffing"); ws.cell(row=r, column=44, value="Staffing-Secured"); r += 2
    write_row(ws, layout, r, "Xray Corp - Staffing", "Wendy", "Staffing", "Xray Corp", "DS70_Secured", 9, 300011, 30011); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Staffing"); r += 2

    ws.cell(row=r, column=1, value="Investments"); ws.cell(row=r, column=44, value="Investments"); r += 2
    write_row(ws, layout, r, "Yankee Corp - Internal", "Priya", "Internal", "Yankee Corp", "DS10_Secured", 12, 400011, 40011); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Investments"); r += 1

    wb.save(path)


def main() -> int:
    fixture = TEST_DIR / "fixtures" / "test6_monthly_view_matches_general.xlsx"
    build_fixture(fixture)

    out_dir = TEST_DIR.parent / "output" / "_test6"
    out_dir.mkdir(parents=True, exist_ok=True)
    rc = sfae_main.main(["--input", str(fixture), "--output-dir", str(out_dir), "--year", "2026"])
    if rc != 0:
        print("FAIL - pipeline did not exit cleanly")
        return 1

    generated = list(out_dir.glob("Sales_and_Forecast_Summary_*.xlsx"))[0]
    recalc_dir = TEST_DIR.parent / "output" / "_test6_recalc"
    recalc_dir.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["soffice", "--headless", "--norestore", "--convert-to", "xlsx:Calc MS Excel 2007 XML",
         "--outdir", str(recalc_dir), str(generated)],
        capture_output=True, timeout=60,
    )
    recalculated = recalc_dir / generated.name

    wb = load_workbook(recalculated, data_only=True)
    ws1, ws2 = wb["Multi-Year Revenue & Margin"], wb["2026 Monthly Performance"]

    def ws1_total(name):
        for r in range(1, ws1.max_row + 1):
            if ws1.cell(row=r, column=1).value == name:
                q = [ws1.cell(row=r, column=c).value for c in (6, 8, 10, 12)]
                return sum(x for x in q if isinstance(x, (int, float)))
        return None

    def ws2_total(name):
        for r in range(1, ws2.max_row + 1):
            if ws2.cell(row=r, column=1).value == name:
                monthly = [ws2.cell(row=r, column=c).value for c in range(3, 27, 2)]
                return sum(x for x in monthly if isinstance(x, (int, float)))
        return None

    failures = []
    for name in ("Victor Corp", "Whiskey Corp", "Xray Corp", "Yankee Corp"):
        v1, v2 = ws1_total(name), ws2_total(name)
        if v1 != v2:
            failures.append(f"{name}: Worksheet1={v1} != Worksheet2={v2}")
        else:
            print(f"PASS - {name}: Worksheet1 == Worksheet2 == {v1}")

    if failures:
        print("FAIL:")
        for f in failures:
            print(" -", f)
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
