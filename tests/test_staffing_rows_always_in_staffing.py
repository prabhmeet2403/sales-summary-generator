"""
tests/test_staffing_rows_always_in_staffing.py
================================================
Verifies requirement 3: a row physically located under the Staffing
heading always appears in the Staffing-Secured section's output, with
its correct revenue -- regardless of what its Sub-Group text says.

Usage:
    python tests/test_staffing_rows_always_in_staffing.py
"""
from __future__ import annotations
import sys
from pathlib import Path

TEST_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TEST_DIR))
sys.path.insert(0, str(TEST_DIR.parent))
from fixture_builder import new_workbook, write_row  # noqa: E402

import main as sfae_main  # noqa: E402
from openpyxl import load_workbook


def build_fixture(path: Path):
    wb, ws, layout = new_workbook()
    r = 4
    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 1"); ws.cell(row=r, column=44, value="Track 1-Secured"); r += 2
    write_row(ws, layout, r, "India Corp - SOW", "Priya", "Project", "India Corp", "DS10_Secured", 1, 111119, 11119); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 1"); r += 2

    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 2"); ws.cell(row=r, column=44, value="Track 2-Secured"); r += 2
    write_row(ws, layout, r, "Juliet Corp - SOW", "Dana", "Staff Aug", "Juliet Corp", "DS20_Secured", 1, 222229, 22229); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 2"); r += 2

    ws.cell(row=r, column=1, value="TOTAL Secured"); r += 2

    ws.cell(row=r, column=1, value="Staffing"); ws.cell(row=r, column=44, value="Staffing-Secured"); r += 2
    write_row(ws, layout, r, "Kilo Corp - Staffing (correct code)", "Wendy", "Staffing", "Kilo Corp", "DS70_Secured", 1, 333339, 33339); r += 1
    write_row(ws, layout, r, "Lima Corp - Staffing (unmapped code)", "Wendy", "Staffing", "Lima Corp", "DS95_Secured", 1, 444449, 44449); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Staffing"); r += 2

    ws.cell(row=r, column=1, value="Investments"); ws.cell(row=r, column=44, value="Investments"); r += 2
    write_row(ws, layout, r, "Mike Corp - Internal", "Priya", "Internal", "Mike Corp", "DS10_Secured", 1, 555559, 55559); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Investments"); r += 1

    wb.save(path)


def main() -> int:
    fixture = TEST_DIR / "fixtures" / "test3_staffing_always_in_staffing.xlsx"
    build_fixture(fixture)

    out_dir = TEST_DIR.parent / "output" / "_test3"
    out_dir.mkdir(parents=True, exist_ok=True)
    rc = sfae_main.main(["--input", str(fixture), "--output-dir", str(out_dir), "--year", "2026"])
    if rc != 0:
        print("FAIL - pipeline did not exit cleanly")
        return 1

    generated = list(out_dir.glob("Sales_and_Forecast_Summary_*.xlsx"))
    wb = load_workbook(generated[0], data_only=False)
    ws2 = wb["2026 Monthly Performance"]

    failures = []
    for expected_name, expected_val in (("Kilo Corp", 333339), ("Lima Corp", 444449)):
        in_staffing = False
        found = False
        for r in range(1, ws2.max_row + 1):
            v = ws2.cell(row=r, column=1).value
            # The fixture's own input heading (line ~37) is "Staffing"
            # -- with the input-heading-propagation fix, that is
            # exactly what the output now shows (not config.py's fixed
            # "Staffing- Secured" title), which is the correct,
            # intended behavior this test should reflect.
            if v == "Staffing":
                in_staffing = True
                continue
            if in_staffing:
                if v and "Subtotal" in str(v):
                    break
                if v == expected_name:
                    found = True
                    total_formula = ws2.cell(row=r, column=27).value
                    monthly = [ws2.cell(row=r, column=c).value for c in range(3, 27, 2)]
                    if not (isinstance(monthly[0], (int, float)) and monthly[0] == expected_val):
                        failures.append(f"{expected_name}: expected Jan={expected_val}, got {monthly[0]!r} (formula={total_formula!r})")
        if not found:
            failures.append(f"{expected_name}: not found inside Staffing- Secured at all")

    if failures:
        print("FAIL:")
        for f in failures:
            print(" -", f)
        return 1
    print("PASS - both Staffing rows (including the unmapped Sub-Group code) show their correct revenue in Staffing")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
