"""
tests/test_track1_never_in_track2.py
=====================================
Verifies requirement 1: a row physically located under the Track 1
heading must never appear inside Track 2's section on either sheet.

Usage:
    python tests/test_track1_never_in_track2.py
"""
from __future__ import annotations
import sys
from pathlib import Path

TEST_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TEST_DIR))
sys.path.insert(0, str(TEST_DIR.parent))
from fixture_builder import new_workbook, write_row  # noqa: E402

import main as sfae_main  # noqa: E402
from openpyxl import load_workbook


def build_fixture(path: Path):
    wb, ws, layout = new_workbook()
    r = 4
    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 1"); ws.cell(row=r, column=44, value="Track 1-Secured"); r += 2
    write_row(ws, layout, r, "Alpha Corp - SOW", "Priya", "Project", "Alpha Corp", "DS10_Secured", 1, 500001, 50001); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 1"); r += 2

    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 2"); ws.cell(row=r, column=44, value="Track 2-Secured"); r += 2
    write_row(ws, layout, r, "Bravo Corp - SOW", "Dana", "Staff Aug", "Bravo Corp", "DS20_Secured", 1, 600002, 60002); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 2"); r += 2

    ws.cell(row=r, column=1, value="TOTAL Secured"); r += 2

    ws.cell(row=r, column=1, value="Staffing"); ws.cell(row=r, column=44, value="Staffing-Secured"); r += 2
    write_row(ws, layout, r, "Charlie Corp - Staffing", "Wendy", "Staffing", "Charlie Corp", "DS70_Secured", 1, 700003, 70003); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Staffing"); r += 2

    ws.cell(row=r, column=1, value="Investments"); ws.cell(row=r, column=44, value="Investments"); r += 2
    write_row(ws, layout, r, "Delta Corp - Internal", "Priya", "Internal", "Delta Corp", "DS10_Secured", 1, 800004, 80004); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Investments"); r += 1

    wb.save(path)


def main() -> int:
    fixture = TEST_DIR / "fixtures" / "test1_track1_never_in_track2.xlsx"
    build_fixture(fixture)

    out_dir = TEST_DIR.parent / "output" / "_test1"
    out_dir.mkdir(parents=True, exist_ok=True)
    rc = sfae_main.main(["--input", str(fixture), "--output-dir", str(out_dir), "--year", "2026"])
    if rc != 0:
        print("FAIL - pipeline did not exit cleanly")
        return 1

    generated = list(out_dir.glob("Sales_and_Forecast_Summary_*.xlsx"))
    wb = load_workbook(generated[0], data_only=False)
    failures = []
    for sheet in ("Multi-Year Revenue & Margin", "2026 Monthly Performance"):
        ws = wb[sheet]
        in_track2 = False
        for r in range(1, ws.max_row + 1):
            v = ws.cell(row=r, column=1).value
            if v == "Solutions and Staff Augmentation (Projects) - Track 2":
                in_track2 = True
                continue
            if in_track2:
                if v and "Subtotal" in str(v):
                    break
                if v == "Alpha Corp":
                    failures.append(f"{sheet}: 'Alpha Corp' (a Track 1 row) found inside Track 2 at row {r}")

    if failures:
        print("FAIL:")
        for f in failures:
            print(" -", f)
        return 1
    print("PASS - Alpha Corp never appears inside Track 2 on either sheet")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
