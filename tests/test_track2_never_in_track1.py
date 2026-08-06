"""
tests/test_track2_never_in_track1.py
=====================================
Verifies requirement 2: a row physically located under the Track 2
heading must never appear inside Track 1's section on either sheet.

Usage:
    python tests/test_track2_never_in_track1.py
"""
from __future__ import annotations
import sys
from pathlib import Path

TEST_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(TEST_DIR))
sys.path.insert(0, str(TEST_DIR.parent))
from fixture_builder import new_workbook, write_row  # noqa: E402

import main as sfae_main  # noqa: E402
from openpyxl import load_workbook


def build_fixture(path: Path):
    wb, ws, layout = new_workbook()
    r = 4
    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 1"); ws.cell(row=r, column=44, value="Track 1-Secured"); r += 2
    write_row(ws, layout, r, "Echo Corp - SOW", "Priya", "Project", "Echo Corp", "DS10_Secured", 1, 111117, 11117); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 1"); r += 2

    ws.cell(row=r, column=1, value="Solutions and Staff Augmentation (Projects) - Track 2"); ws.cell(row=r, column=44, value="Track 2-Secured"); r += 2
    write_row(ws, layout, r, "Foxtrot Corp - SOW", "Dana", "Staff Aug", "Foxtrot Corp", "DS20_Secured", 1, 222227, 22227); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Track 2"); r += 2

    ws.cell(row=r, column=1, value="TOTAL Secured"); r += 2

    ws.cell(row=r, column=1, value="Staffing"); ws.cell(row=r, column=44, value="Staffing-Secured"); r += 2
    write_row(ws, layout, r, "Golf Corp - Staffing", "Wendy", "Staffing", "Golf Corp", "DS70_Secured", 1, 333337, 33337); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Staffing"); r += 2

    ws.cell(row=r, column=1, value="Investments"); ws.cell(row=r, column=44, value="Investments"); r += 2
    write_row(ws, layout, r, "Hotel Corp - Internal", "Priya", "Internal", "Hotel Corp", "DS10_Secured", 1, 444447, 44447); r += 1
    ws.cell(row=r, column=1, value="Sub-total : Investments"); r += 1

    wb.save(path)


def main() -> int:
    fixture = TEST_DIR / "fixtures" / "test2_track2_never_in_track1.xlsx"
    build_fixture(fixture)

    out_dir = TEST_DIR.parent / "output" / "_test2"
    out_dir.mkdir(parents=True, exist_ok=True)
    rc = sfae_main.main(["--input", str(fixture), "--output-dir", str(out_dir), "--year", "2026"])
    if rc != 0:
        print("FAIL - pipeline did not exit cleanly")
        return 1

    generated = list(out_dir.glob("Sales_and_Forecast_Summary_*.xlsx"))
    wb = load_workbook(generated[0], data_only=False)
    failures = []
    for sheet in ("Multi-Year Revenue & Margin", "2026 Monthly Performance"):
        ws = wb[sheet]
        in_track1 = False
        for r in range(1, ws.max_row + 1):
            v = ws.cell(row=r, column=1).value
            if v == "Solutions and Staff Augmentation (Projects) - Track 1":
                in_track1 = True
                continue
            if in_track1:
                if v and "Subtotal" in str(v):
                    break
                if v == "Foxtrot Corp":
                    failures.append(f"{sheet}: 'Foxtrot Corp' (a Track 2 row) found inside Track 1 at row {r}")

    if failures:
        print("FAIL:")
        for f in failures:
            print(" -", f)
        return 1
    print("PASS - Foxtrot Corp never appears inside Track 1 on either sheet")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
